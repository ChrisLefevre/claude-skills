#!/usr/bin/env python3
"""Lecture et analyse d'un fichier Excel de suivi de projets.

Detecte les colonnes par leur en-tete, classe les projets par statut,
calcule les echeances et sort un etat lisible pour rediger un brief.

Usage:
    python3 lire_suivi.py "suivi-projets.xlsx"
    python3 lire_suivi.py "suivi-projets.xlsx" --jours 10 --date 2026-08-24
    python3 lire_suivi.py "suivi-projets.xlsx" --onglet Projets --journal "Suivi quotidien"
    python3 lire_suivi.py "suivi-projets.xlsx" --json
"""

import argparse
import datetime as dt
import fnmatch
import json
import re
import sys
import unicodedata

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl requis : pip install openpyxl --break-system-packages")


# --- correspondance des colonnes -------------------------------------------

ALIAS = {
    "id": ["id", "n", "no", "numero", "ref", "reference"],
    "priorite": ["priorite", "priority", "importance"],
    "statut": ["statut", "status", "etat"],
    "domaine": ["domaine", "categorie", "axe", "domain", "category", "client"],
    "projet": ["projet", "intitule", "nom", "tache", "libelle", "project", "task"],
    "debut_prevu": ["debut prevu", "date debut", "start", "debut planifie"],
    "fin_prevue": ["fin prevue", "echeance", "deadline", "date fin", "due", "butoir", "a finir pour"],
    "debut_reel": ["debut reel", "demarre le", "actual start"],
    "fin_reelle": ["fin reelle", "termine le", "cloture le", "actual end", "date de cloture"],
    "jours_estimes": ["jours estimes", "charge prevue j", "charge prevue", "charge estimee", "estime", "estimation", "estimated"],
    "jours_reels": ["jours reels", "temps passe j", "temps passe", "charge reelle", "realise", "actual days"],
    "ecart": ["ecart", "ecart (j)", "variance"],
    "avancement": ["avancement", "progression", "progress", "%", "pourcentage"],
    "prochaine_etape": ["prochaine etape", "next step", "action", "prochaine action"],
    "equipe": ["team", "responsable", "owner", "assigne", "porteur"],
    "kpi": ["kpi", "kpi lie", "indicateur"],
    "lien": ["lien", "lien odoo", "url", "link"],
    "notes": ["notes", "details", "commentaires", "notes / details", "remarques"],
}

FAMILLES_STATUT = {
    "en cours": "en_cours",
    "in progress": "en_cours",
    "demarre": "en_cours",
    "prevu": "prevu",
    "planned": "prevu",
    "a faire": "prevu",
    "todo": "prevu",
    "bloque": "bloque",
    "blocked": "bloque",
    "en attente": "bloque",
    "termine": "termine",
    "done": "termine",
    "cloture": "termine",
    "fini": "termine",
    "annule": "annule",
    "cancelled": "annule",
    "abandonne": "annule",
    "backlog": "backlog",
    "idee": "backlog",
}

FAMILLES_CLOSES = {"termine", "annule"}


def normaliser(texte):
    """Minuscules, sans accents, sans ponctuation superflue."""
    if texte is None:
        return ""
    texte = str(texte).strip().lower()
    texte = unicodedata.normalize("NFD", texte)
    texte = "".join(c for c in texte if unicodedata.category(c) != "Mn")
    texte = re.sub(r"[^\w%/ ]+", " ", texte)
    return re.sub(r"\s+", " ", texte).strip()


def trouver_ligne_entetes(ws, max_scan=10):
    """Repere la ligne d'en-tetes : celle qui reconnait le plus de colonnes."""
    meilleur, score_max = 1, 0
    for r in range(1, min(max_scan, ws.max_row) + 1):
        valeurs = [normaliser(c.value) for c in ws[r]]
        score = sum(1 for v in valeurs if v and any(v in a or a in v for al in ALIAS.values() for a in al))
        if score > score_max:
            meilleur, score_max = r, score
    return meilleur


def mapper_colonnes(ws, ligne_entetes, surcharges=None):
    """Associe chaque role a un index de colonne."""
    surcharges = surcharges or {}
    entetes = {}
    for c in ws[ligne_entetes]:
        n = normaliser(c.value)
        if n:
            entetes[n] = c.column

    mapping = {}
    for role, alias in ALIAS.items():
        if role in surcharges:
            cible = normaliser(surcharges[role])
            if cible in entetes:
                mapping[role] = entetes[cible]
                continue
        for a in alias:
            if a in entetes:
                mapping[role] = entetes[a]
                break
        else:
            for tete, col in entetes.items():
                if any(tete.startswith(a) or a in tete for a in alias):
                    mapping[role] = col
                    break
    return mapping


def famille_statut(valeur):
    n = normaliser(valeur)
    n = re.sub(r"^\d+[\.\)]?\s*", "", n)
    for cle, famille in FAMILLES_STATUT.items():
        if cle in n:
            return famille
    return "autre" if n else "vide"


def en_date(valeur):
    if isinstance(valeur, dt.datetime):
        return valeur.date()
    if isinstance(valeur, dt.date):
        return valeur
    if isinstance(valeur, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y"):
            try:
                return dt.datetime.strptime(valeur.strip(), fmt).date()
            except ValueError:
                continue
    return None


def en_pourcent(valeur):
    if valeur is None:
        return None
    if isinstance(valeur, (int, float)):
        return round(valeur * 100) if valeur <= 1 else round(valeur)
    m = re.search(r"(\d+(?:[.,]\d+)?)", str(valeur))
    return round(float(m.group(1).replace(",", "."))) if m else None


def lire_projets(ws, mapping, ligne_entetes):
    projets = []
    for r in range(ligne_entetes + 1, ws.max_row + 1):
        def val(role):
            col = mapping.get(role)
            return ws.cell(row=r, column=col).value if col else None

        identifiant = val("id")
        libelle = val("projet")
        # ligne vide ou ligne de legende (texte long sans identifiant)
        if identifiant is None and not isinstance(identifiant, (int, float)):
            continue
        if libelle is None and val("statut") is None:
            continue

        projets.append({
            "ligne": r,
            "id": identifiant,
            "priorite": val("priorite"),
            "statut": val("statut"),
            "famille": famille_statut(val("statut")),
            "domaine": val("domaine"),
            "projet": str(libelle) if libelle else "",
            "debut_prevu": en_date(val("debut_prevu")),
            "fin_prevue": en_date(val("fin_prevue")),
            "debut_reel": en_date(val("debut_reel")),
            "fin_reelle": en_date(val("fin_reelle")),
            "jours_estimes": val("jours_estimes"),
            "jours_reels": val("jours_reels"),
            "avancement": en_pourcent(val("avancement")),
            "prochaine_etape": val("prochaine_etape"),
            "equipe": val("equipe"),
            "notes": val("notes"),
        })
    return projets


def niveau_priorite(p):
    if p is None:
        return 0
    s = str(p)
    etoiles = s.count("⭐") or s.count("*")
    if etoiles:
        return etoiles
    m = re.search(r"\d", s)
    return int(m.group()) if m else 0


def analyser(projets, aujourdhui, horizon):
    actifs = [p for p in projets if p["famille"] not in FAMILLES_CLOSES]
    for p in actifs:
        p["jours_restants"] = (p["fin_prevue"] - aujourdhui).days if p["fin_prevue"] else None
        p["etoiles"] = niveau_priorite(p["priorite"])

    depassees = sorted(
        [p for p in actifs if p["jours_restants"] is not None and p["jours_restants"] < 0],
        key=lambda p: p["jours_restants"])
    proches = sorted(
        [p for p in actifs if p["jours_restants"] is not None and 0 <= p["jours_restants"] <= horizon],
        key=lambda p: p["jours_restants"])
    bloques = [p for p in actifs if p["famille"] == "bloque"]
    demarrages = [p for p in actifs
                  if p["famille"] == "prevu" and p["debut_prevu"] and p["debut_prevu"] <= aujourdhui
                  and not p["debut_reel"]]
    prioritaires = sorted(
        [p for p in actifs if p["etoiles"] >= 3 and p["famille"] == "en_cours"],
        key=lambda p: (-(p["avancement"] or 0)))
    sans_echeance = [p for p in actifs if not p["fin_prevue"] and p["famille"] in ("en_cours", "bloque")]

    # charge cumulee sur la fenetre a venir
    fin_fenetre = aujourdhui + dt.timedelta(days=horizon * 2)
    charge = sum(p["jours_estimes"] or 0 for p in actifs
                 if isinstance(p["jours_estimes"], (int, float))
                 and p["fin_prevue"] and aujourdhui <= p["fin_prevue"] <= fin_fenetre)

    incoherences = []
    for p in actifs:
        if p["debut_prevu"] and p["fin_prevue"] and p["fin_prevue"] < p["debut_prevu"]:
            incoherences.append(f"#{p['id']} fin prevue anterieure au debut prevu")
        if p["avancement"] == 100 and p["famille"] == "en_cours":
            incoherences.append(f"#{p['id']} a 100 % mais toujours en cours")
        if p["fin_reelle"] and p["famille"] not in FAMILLES_CLOSES:
            incoherences.append(f"#{p['id']} a une fin reelle mais n'est pas cloture")

    return {
        "actifs": actifs,
        "depassees": depassees,
        "proches": proches,
        "bloques": bloques,
        "demarrages": demarrages,
        "prioritaires": prioritaires,
        "sans_echeance": sans_echeance,
        "charge_fenetre": round(charge, 2),
        "fin_fenetre": fin_fenetre,
        "incoherences": incoherences,
    }


def lire_journal(wb, nom_onglet, nb):
    if not nom_onglet or nom_onglet not in wb.sheetnames:
        return []
    ws = wb[nom_onglet]
    lignes = []
    for r in range(1, ws.max_row + 1):
        valeurs = [ws.cell(row=r, column=c).value for c in range(1, min(ws.max_column, 6) + 1)]
        if any(v is not None for v in valeurs):
            d = en_date(valeurs[0])
            if d or any(isinstance(v, str) and len(str(v)) > 15 for v in valeurs[1:]):
                lignes.append({"date": d, "valeurs": [v for v in valeurs if v is not None]})
    return lignes[-nb:]


def ligne_projet(p):
    ech = p["fin_prevue"].strftime("%d/%m") if p["fin_prevue"] else "-"
    jr = p.get("jours_restants")
    marqueur = f" (J{jr:+d})" if jr is not None else ""
    av = f"{p['avancement']}%" if p["avancement"] is not None else "-"
    prio = str(p["priorite"] or "").strip() or "-"
    etape = str(p["prochaine_etape"] or "").strip().replace("\n", " ")
    if len(etape) > 110:
        etape = etape[:107] + "..."
    return (f"  L{p['ligne']} #{p['id']} [{prio}] {p['projet'][:58]}\n"
            f"      statut={p['statut']} | echeance={ech}{marqueur} | avancement={av} | "
            f"estime={p['jours_estimes']} reel={p['jours_reels']}\n"
            f"      etape: {etape or '-'}")


def rapport_texte(a, journal, aujourdhui, horizon):
    out = []
    out.append(f"ETAT DU SUIVI au {aujourdhui.strftime('%d/%m/%Y')}")
    out.append(f"{len(a['actifs'])} projets actifs, horizon d'echeance {horizon} jours\n")

    sections = [
        ("ECHEANCES DEPASSEES", a["depassees"]),
        (f"ECHEANCES SOUS {horizon} JOURS", a["proches"]),
        ("PROJETS BLOQUES", a["bloques"]),
        ("DEMARRAGES ATTENDUS (prevus, non demarres)", a["demarrages"]),
        ("PRIORITAIRES EN COURS", a["prioritaires"]),
        ("ACTIFS SANS ECHEANCE", a["sans_echeance"]),
    ]
    for titre, items in sections:
        out.append(f"== {titre} ({len(items)})")
        out.extend(ligne_projet(p) for p in items) if items else out.append("  aucun")
        out.append("")

    out.append(f"== CHARGE estimee sur les projets a echeance avant le "
               f"{a['fin_fenetre'].strftime('%d/%m')} : {a['charge_fenetre']} jours")
    out.append("")

    if a["incoherences"]:
        out.append("== INCOHERENCES A SIGNALER (ne pas corriger seul)")
        out.extend(f"  - {i}" for i in a["incoherences"])
        out.append("")

    if journal:
        out.append(f"== DERNIERES ENTREES DU JOURNAL ({len(journal)})")
        for e in journal:
            d = e["date"].strftime("%d/%m") if e["date"] else "?"
            corps = " | ".join(str(v)[:120] for v in e["valeurs"][1:])
            out.append(f"  {d} : {corps}")
    return "\n".join(out)


def main():
    ap = argparse.ArgumentParser(description="Analyse d'un fichier de suivi de projets")
    ap.add_argument("fichier")
    ap.add_argument("--onglet", default=None, help="onglet des projets (defaut : detection auto)")
    ap.add_argument("--journal", default=None, help="onglet du journal quotidien")
    ap.add_argument("--jours", type=int, default=7, help="horizon d'echeance en jours")
    ap.add_argument("--date", default=None, help="date de reference AAAA-MM-JJ (defaut : aujourd'hui)")
    ap.add_argument("--entrees-journal", type=int, default=12)
    ap.add_argument("--tous", action="store_true", help="inclure les projets clos")
    ap.add_argument("--json", action="store_true", help="sortie JSON")
    ap.add_argument("--colonnes", default=None, help="surcharges JSON, ex: '{\"fin_prevue\":\"Butoir\"}'")
    args = ap.parse_args()

    aujourdhui = dt.datetime.strptime(args.date, "%Y-%m-%d").date() if args.date else dt.date.today()
    wb = openpyxl.load_workbook(args.fichier, data_only=True)

    nom = args.onglet
    if not nom:
        for candidat in wb.sheetnames:
            if normaliser(candidat) in ("projets", "projects", "suivi projets", "projets et taches"):
                nom = candidat
                break
        else:
            nom = wb.sheetnames[0]
    if nom not in wb.sheetnames:
        sys.exit(f"onglet '{nom}' absent. Onglets : {wb.sheetnames}")

    nom_journal = args.journal
    if nom_journal is None:
        for candidat in wb.sheetnames:
            if "quotidien" in normaliser(candidat) or "journal" in normaliser(candidat):
                nom_journal = candidat
                break

    ws = wb[nom]
    ligne_entetes = trouver_ligne_entetes(ws)
    surcharges = json.loads(args.colonnes) if args.colonnes else None
    mapping = mapper_colonnes(ws, ligne_entetes, surcharges)

    manquantes = [r for r in ("id", "statut", "projet") if r not in mapping]
    if manquantes:
        sys.exit(f"colonnes non trouvees : {manquantes}. Utiliser --colonnes pour les declarer.")

    projets = lire_projets(ws, mapping, ligne_entetes)
    if args.tous:
        analyse = analyser(projets, aujourdhui, args.jours)
        analyse["actifs"] = projets
    else:
        analyse = analyser(projets, aujourdhui, args.jours)
    journal = lire_journal(wb, nom_journal, args.entrees_journal)

    if args.json:
        def serialiser(o):
            return o.isoformat() if isinstance(o, (dt.date, dt.datetime)) else str(o)
        print(json.dumps({
            "date": aujourdhui.isoformat(),
            "onglet": nom,
            "onglet_journal": nom_journal,
            "colonnes": {k: v for k, v in mapping.items()},
            "analyse": {k: v for k, v in analyse.items() if k != "actifs"},
            "projets_actifs": analyse["actifs"],
            "journal": journal,
        }, default=serialiser, ensure_ascii=False, indent=2))
    else:
        print(rapport_texte(analyse, journal, aujourdhui, args.jours))


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""Application en dur des couleurs de statut dans un fichier de suivi.

La mise en forme conditionnelle n'est pas rendue par tous les visualiseurs
(apercu rapide macOS notamment). Ce script pose les couleurs directement
dans les cellules : statut en couleur franche, ligne en teinte pastel,
lignes closes grisees.

Usage:
    python3 appliquer_statut.py "suivi-projets.xlsx"
    python3 appliquer_statut.py "suivi-projets.xlsx" --onglet Projets --lignes 8,65,69
    python3 appliquer_statut.py "suivi-projets.xlsx" --palette palette.json --sortie copie.xlsx
"""

import argparse
import json
import re
import sys
import unicodedata

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
except ImportError:
    sys.exit("openpyxl requis : pip install openpyxl --break-system-packages")


PALETTE = {
    "en_cours": "548235",
    "prevu": "2E75B6",
    "bloque": "C00000",
    "backlog": "BF8F00",
    "termine": "A6A6A6",
    "annule": "D9D9D9",
}

FAMILLES = {
    "en cours": "en_cours", "in progress": "en_cours", "demarre": "en_cours",
    "prevu": "prevu", "planned": "prevu", "a faire": "prevu", "todo": "prevu",
    "bloque": "bloque", "blocked": "bloque", "en attente": "bloque",
    "termine": "termine", "done": "termine", "cloture": "termine", "fini": "termine",
    "annule": "annule", "cancelled": "annule", "abandonne": "annule",
    "backlog": "backlog", "idee": "backlog",
}

CLOSES = {"termine", "annule"}
GRIS_TEXTE = "808080"


def normaliser(t):
    if t is None:
        return ""
    t = str(t).strip().lower()
    t = unicodedata.normalize("NFD", t)
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", re.sub(r"[^\w% ]+", " ", t)).strip()


def famille(valeur):
    n = re.sub(r"^\d+[\.\)]?\s*", "", normaliser(valeur))
    for cle, f in FAMILLES.items():
        if cle in n:
            return f
    return None


def pastel(hexa, force=0.86):
    """Melange la couleur avec du blanc. force=0.86 donne une teinte tres claire."""
    r, g, b = int(hexa[0:2], 16), int(hexa[2:4], 16), int(hexa[4:6], 16)
    melange = lambda c: int(c + (255 - c) * force)
    return f"{melange(r):02X}{melange(g):02X}{melange(b):02X}"


def trouver_colonne_statut(ws, max_scan=10):
    for r in range(1, min(max_scan, ws.max_row) + 1):
        for c in ws[r]:
            if normaliser(c.value) in ("statut", "status", "etat"):
                return c.column, r
    return None, None


def derniere_colonne_utile(ws, ligne_entetes):
    derniere = 1
    for c in ws[ligne_entetes]:
        if c.value not in (None, ""):
            derniere = c.column
    return derniere


def appliquer(chemin, nom_onglet, lignes_cibles, palette, sortie, force_pastel):
    wb = openpyxl.load_workbook(chemin)
    if nom_onglet and nom_onglet not in wb.sheetnames:
        sys.exit(f"onglet '{nom_onglet}' absent. Onglets : {wb.sheetnames}")
    ws = wb[nom_onglet] if nom_onglet else wb[wb.sheetnames[0]]

    col_statut, ligne_entetes = trouver_colonne_statut(ws)
    if not col_statut:
        sys.exit("colonne de statut introuvable (intitules cherches : Statut, Status, Etat)")

    col_max = derniere_colonne_utile(ws, ligne_entetes)
    lignes = lignes_cibles or range(ligne_entetes + 1, ws.max_row + 1)

    traitees, ignorees = 0, 0
    for r in lignes:
        cellule = ws.cell(row=r, column=col_statut)
        f = famille(cellule.value)
        if not f or f not in palette:
            ignorees += 1
            continue

        franc = palette[f]
        clair = pastel(franc, force_pastel)

        # ligne entiere en teinte pastel
        for c in range(1, col_max + 1):
            cel = ws.cell(row=r, column=c)
            cel.fill = PatternFill("solid", fgColor=clair)
            if f in CLOSES:
                ancienne = cel.font
                cel.font = Font(name=ancienne.name, size=ancienne.size, bold=ancienne.bold,
                                italic=ancienne.italic, color=GRIS_TEXTE)

        # cellule de statut en couleur franche, texte blanc gras
        cellule.fill = PatternFill("solid", fgColor=franc)
        cellule.font = Font(name=cellule.font.name, size=cellule.font.size,
                            bold=True, color="FFFFFF")
        traitees += 1

    destination = sortie or chemin
    wb.save(destination)
    print(f"{traitees} lignes coloriees, {ignorees} ignorees (statut vide ou inconnu)")
    print(f"onglet : {ws.title} | colonne statut : {col_statut} | fichier : {destination}")


def main():
    ap = argparse.ArgumentParser(description="Couleurs de statut en dur")
    ap.add_argument("fichier")
    ap.add_argument("--onglet", default=None)
    ap.add_argument("--lignes", default=None, help="lignes a traiter, ex: 8,65,69 (defaut : toutes)")
    ap.add_argument("--palette", default=None, help="fichier JSON de palette personnalisee")
    ap.add_argument("--sortie", default=None, help="fichier de sortie (defaut : modification sur place)")
    ap.add_argument("--pastel", type=float, default=0.86, help="intensite de l'eclaircissement, 0 a 1")
    args = ap.parse_args()

    palette = dict(PALETTE)
    if args.palette:
        with open(args.palette, encoding="utf-8") as f:
            palette.update(json.load(f))

    lignes = [int(x) for x in args.lignes.split(",")] if args.lignes else None
    appliquer(args.fichier, args.onglet, lignes, palette, args.sortie, args.pastel)


if __name__ == "__main__":
    main()
